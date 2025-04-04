import os
from flask import Blueprint, send_file, abort, current_app
from flask_login import login_required
import tempfile
from pdf2image import convert_from_path
from docx import Document
from openpyxl import load_workbook
from ..utils import logger
from werkzeug.utils import secure_filename

preview_bp = Blueprint('preview', __name__)

@preview_bp.route('/preview/<path:filename>')
@login_required
def preview_document(filename):
    try:
        safe_filename = secure_filename(filename)
        file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], safe_filename)
        
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            abort(404)
            
        file_ext = os.path.splitext(filename)[1].lower()
        
        # PDF files
        if file_ext == '.pdf':
            return send_file(file_path)
            
        # Word documents
        elif file_ext in ['.doc', '.docx']:
            doc = Document(file_path)
            text_content = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
            
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.txt')
            temp_file.write(text_content.encode('utf-8'))
            temp_file.close()
            
            return send_file(temp_file.name, mimetype='text/plain')
            
        # Excel files
        elif file_ext in ['.xls', '.xlsx']:
            wb = load_workbook(filename=file_path, read_only=True)
            sheet = wb.active
            text_content = '\n'.join([str(row) for row in sheet.values])
            
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.txt')
            temp_file.write(text_content.encode('utf-8'))
            temp_file.close()
            
            return send_file(temp_file.name, mimetype='text/plain')
            
        # Text files
        elif file_ext == '.txt':
            return send_file(file_path, mimetype='text/plain')
            
        else:
            logger.error(f"Unsupported file type: {file_ext}")
            abort(400, description="Unsupported file type")
            
    except Exception as e:
        logger.error(f"Error previewing document: {str(e)}")
        abort(500, description="Error generating preview")
