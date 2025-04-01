import os
import base64
import mimetypes
from io import BytesIO
from flask import Blueprint, render_template, redirect, url_for, flash
from flask_login import login_required
from werkzeug.utils import secure_filename
from pdf2image import convert_from_bytes
from docx import Document as DocxDocument
from openpyxl import load_workbook
import filetype
from app import logger
from config import Config

preview_bp = Blueprint('preview', __name__)

@preview_bp.route('/preview/<filename>')
@login_required
def preview_file(filename):
    safe_filename = secure_filename(filename)
    file_path = os.path.join(Config.UPLOAD_FOLDER, safe_filename)

    if not os.path.exists(file_path):
        flash('File not found', 'error')
        return redirect(url_for('documents.documents'))

    mime_type = None
    try:
        kind = filetype.guess(file_path)
        if kind:
            mime_type = kind.mime
        if not mime_type:
            guess, _ = mimetypes.guess_type(file_path)
            mime_type = guess or 'application/octet-stream'
    except Exception as e:
        logger.error("Error detecting file type in preview: %s", str(e))
        mime_type = 'application/octet-stream'

    try:
        if mime_type == 'application/pdf':
            with open(file_path, 'rb') as f:
                pdf_bytes = f.read()
            images = convert_from_bytes(pdf_bytes, first_page=1, last_page=1, dpi=100, poppler_path=Config.POPPLER_PATH)
            buffered = BytesIO()
            images[0].save(buffered, format="JPEG")
            img_str = base64.b64encode(buffered.getvalue()).decode()
            return render_template('preview.html', content=f'data:image/jpeg;base64,{img_str}', filename=safe_filename)

        elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            doc = DocxDocument(file_path)
            full_text = [para.text for para in doc.paragraphs if para.text.strip()]
            preview_content = '<br>'.join(full_text[:20])
            return render_template('preview.html', content=preview_content, filename=safe_filename)

        elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            wb = load_workbook(file_path, read_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            limited_rows = rows[:10]  # İlk 10 satırı gösterelim

            table_html = "<table class='table table-bordered table-sm'>"
            for row in limited_rows:
                table_html += "<tr>" + "".join(f"<td>{cell if cell is not None else ''}</td>" for cell in row) + "</tr>"
            table_html += "</table>"

            return render_template('preview.html', content=table_html, filename=safe_filename)

        elif mime_type in ['image/jpeg', 'image/png']:
            with open(file_path, 'rb') as f:
                image_data = f.read()
            img_str = base64.b64encode(image_data).decode()
            return render_template('preview.html', content=f'data:{mime_type};base64,{img_str}', filename=safe_filename)

        elif mime_type.startswith('application/') or mime_type.startswith('image/'):
            flash('Bu dosya türü için önizleme desteklenmiyor.', 'info')
            return redirect(url_for('documents.documents'))

        else:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read(2000)
            return render_template('preview.html', content=content, filename=safe_filename)

    except Exception as e:
        logger.error("Preview error: %s", str(e))
        flash('Dosya önizleme hatası', 'error')
        return redirect(url_for('documents.documents'))
